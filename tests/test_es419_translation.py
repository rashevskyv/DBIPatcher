from __future__ import annotations

import csv
import re
import struct
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
ES_CSV = ROOT / "translations" / "es.csv"
ES419_CSV = ROOT / "translations" / "es419.csv"

sys.path.insert(0, str(ROOT))

from src.core.validator import Validator  # noqa: E402
from src.core.text_utils import tokenize  # noqa: E402


EXPECTED_HEADERS = ["original", "translation"]
ALLOWED_CYRILLIC_TRANSLATIONS = {
    r"\x1b[31;1mГ\x1b[37;1m",
    r"\x1b[32;1mШ\x1b[37;1m",
}
MOJIBAKE_MARKERS = ("\ufffd", "Ã", "Â", "â€", "ðŸ", "Ð")


def load_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


class Es419TranslationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.headers, cls.rows = load_csv(ES419_CSV)
        _, cls.es_rows = load_csv(ES_CSV)
        cls.validator = Validator(ROOT / "data" / "blocks.json")

    def test_csv_schema_and_completeness(self) -> None:
        self.assertEqual(self.headers, EXPECTED_HEADERS)
        self.assertGreater(len(self.rows), 0)

        empty_rows = [
            index
            for index, row in enumerate(self.rows, start=2)
            if not row["original"].strip() or not row["translation"].strip()
        ]
        self.assertEqual(empty_rows, [], f"Empty cells in CSV rows: {empty_rows}")

    def test_original_strings_are_unique(self) -> None:
        originals = [row["original"] for row in self.rows]
        self.assertEqual(len(originals), len(set(originals)))

    def test_es_and_es419_use_the_same_source_strings(self) -> None:
        es_originals = [row["original"] for row in self.es_rows]
        es419_originals = [row["original"] for row in self.rows]
        self.assertEqual(es419_originals, es_originals)

    def test_all_rows_pass_structural_validation(self) -> None:
        errors: list[str] = []
        for index, row in enumerate(self.rows, start=2):
            row_errors = self.validator.validate_row(
                row["original"], row["translation"], "es419"
            )
            errors.extend(f"row {index}: {error}" for error in row_errors)
        self.assertEqual(errors, [], "\n".join(errors))

    def test_no_unexpected_cyrillic_remains(self) -> None:
        unexpected = [
            (index, row["translation"])
            for index, row in enumerate(self.rows, start=2)
            if re.search(r"[\u0400-\u04ff]", row["translation"])
            and row["translation"] not in ALLOWED_CYRILLIC_TRANSLATIONS
        ]
        self.assertEqual(unexpected, [])

    def test_no_common_mojibake_markers(self) -> None:
        broken = [
            (index, marker)
            for index, row in enumerate(self.rows, start=2)
            for marker in MOJIBAKE_MARKERS
            if marker in row["translation"]
        ]
        self.assertEqual(broken, [])

    def test_preferred_es419_terminology(self) -> None:
        forbidden = {
            r"\bajustes?\b": "Configuración",
            r"\bpulsa\b": "Presiona",
            r"\bcopias? de seguridad\b": "Respaldo",
            r"\bpartidas? guardadas?\b": "Datos de guardado",
            r"\bsticks?\b": "Palancas",
            r"\blanzar\b": "Iniciar",
            r"\bborrando\b": "Eliminando",
        }
        errors = []
        for index, row in enumerate(self.rows, start=2):
            for pattern, preferred in forbidden.items():
                if re.search(pattern, row["translation"], re.IGNORECASE):
                    errors.append(
                        f"row {index}: use {preferred!r} instead of "
                        f"{row['translation']!r}"
                    )
        self.assertEqual(errors, [], "\n".join(errors))

    def test_language_code_is_es419(self) -> None:
        translations_by_source = {
            row["original"]: row["translation"] for row in self.rows
        }
        self.assertEqual(translations_by_source.get("ru"), "es419")

    def test_master_dictionary_matches_es419_csv(self) -> None:
        workbook = load_workbook(
            ROOT / "data" / "dictionary.xlsx", read_only=True, data_only=False
        )
        sheet = workbook["Translations"]
        workbook_rows = sheet.iter_rows(values_only=True)
        headers = list(next(workbook_rows))
        original_column = headers.index("Original")
        es419_column = headers.index("es419")
        dictionary = {}
        for row in workbook_rows:
            original = row[original_column]
            if original:
                dictionary[str(original)] = str(row[es419_column] or "")
        workbook.close()

        csv_values = {
            tokenize(row["original"]): tokenize(row["translation"])
            for row in self.rows
        }
        self.assertEqual(csv_values, dictionary)

    def test_binary_builder_produces_a_valid_table(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "translation_es419.bin"
            subprocess.run(
                [
                    sys.executable,
                    str(ROOT / "scripts" / "build_translation_bin.py"),
                    str(ES419_CSV),
                    "-o",
                    str(output),
                ],
                cwd=ROOT,
                check=True,
                capture_output=True,
                text=True,
            )
            data = output.read_bytes()

        magic, version, count, entry_size, ru_offset, tr_offset, reserved = (
            struct.unpack_from("<8sIIIIII", data)
        )
        self.assertEqual(magic, b"DBITRNS\x00")
        self.assertEqual(version, 2)
        self.assertGreater(count, 0)
        self.assertGreater(entry_size, tr_offset)
        self.assertEqual(ru_offset, 4)
        self.assertEqual(reserved, 0)
        self.assertEqual(len(data), 0x20 + count * entry_size)


if __name__ == "__main__":
    unittest.main()
