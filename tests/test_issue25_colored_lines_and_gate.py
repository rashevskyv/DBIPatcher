from __future__ import annotations
import csv, json, sys, unittest
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.core.validator import Validator

def simulate_pfxsfx(data_bytes: bytes, table: dict[bytes, bytes], gate: int = 128) -> bytes:
    orig = data_bytes
    idx = 0
    while idx < len(data_bytes) and data_bytes[idx] < 0x20:
        idx += 1
    skipped = data_bytes[:idx]
    core = data_bytes[idx:]
    if core in table:
        return table[core]
    if len(core) > gate or len(core) < 2:
        return orig
    if not any(b >= 0x80 for b in core):
        return orig
    pfx_len, pfx_val = 0, b''
    for l in range(len(core), 1, -1):
        cand = core[:l]
        if cand in table:
            pfx_len, pfx_val = l, table[cand]
            break
    sfx_len, sfx_val = 0, b''
    for l in range(len(core), 1, -1):
        cand = core[len(core) - l :]
        if cand in table:
            sfx_len, sfx_val = l, table[cand]
            break
    if pfx_len == 0 and sfx_len == 0:
        return orig
    if pfx_len + sfx_len > len(core):
        sfx_len, sfx_val = 0, b''
    mid = core[pfx_len : len(core) - sfx_len]
    return skipped + pfx_val + mid + sfx_val

class Issue25ColoredLinesAndGateTests(unittest.TestCase):
    def setUp(self) -> None:
        self.dict_path = ROOT / 'data' / 'dictionary.xlsx'
        self.lang_path = ROOT / 'data' / 'languages.json'
        self.translations_dir = ROOT / 'translations'
        self.output_dir = ROOT / 'output'
        with open(self.lang_path, 'r', encoding='utf-8') as f:
            self.languages = json.load(f)
        self.expected_keys = [
            '[32;1mУстановка игры завершена[[ESC]][37;1m',
            '[32;1mТикет исправлен[[ESC]][37;1m',
            '[32;1m[ПЕРЕДАЧА OK][[ESC]][37;1m',
            '[[ESC]][32;1m[ПЕРЕДАЧА OK][[ESC]][37;1m',
            '[[ESC]][31;1m[ПЕРЕДАЧА ПРЕРВАНА][[ESC]][37;1m',
            '[ПОДПИСЬ: Поддельная][[ESC]][37;1m',
            '[ПОДПИСЬ: OK][[ESC]][37;1m',
            '[ПОДПИСЬ: XCI➡NSP][[ESC]][37;1m',
            '[[ESC]][32;1m[ПЕРЕДАЧА OK][[ESC]][32;1m[ПОДПИСЬ: OK][[ESC]][37;1m',
            '[[ESC]][32;1m[ПЕРЕДАЧА OK][[ESC]][38;2;255;128;0m[ПОДПИСЬ: Поддельная][[ESC]][37;1m',
            '[[ESC]][32;1m[ПЕРЕДАЧА OK][[ESC]][32;1m[ПОДПИСЬ: XCI➡NSP][[ESC]][37;1m',
        ]

    def test_pfxsfx_simulation_exact_hit_colored_lines(self) -> None:
        ESC = b'\x1b'
        line_finish = ESC + b'[32;1m' + 'Установка игры завершена'.encode('utf-8') + ESC + b'[37;1m'
        line_ticket = ESC + b'[32;1m' + 'Тикет исправлен'.encode('utf-8') + ESC + b'[37;1m'
        lookup_table = {
            b'[32;1m' + 'Установка игры завершена'.encode('utf-8') + ESC + b'[37;1m': ESC + b'[32;1m' + 'Встановлення гри завершено'.encode('utf-8') + ESC + b'[37;1m',
            b'[32;1m' + 'Тикет исправлен'.encode('utf-8') + ESC + b'[37;1m': ESC + b'[32;1m' + 'Тікет виправлено'.encode('utf-8') + ESC + b'[37;1m',
        }
        res_finish = simulate_pfxsfx(line_finish, lookup_table)
        self.assertEqual(res_finish, ESC + b'[32;1m' + 'Встановлення гри завершено'.encode('utf-8') + ESC + b'[37;1m')
        res_ticket = simulate_pfxsfx(line_ticket, lookup_table)
        self.assertEqual(res_ticket, ESC + b'[32;1m' + 'Тікет виправлено'.encode('utf-8') + ESC + b'[37;1m')

    def test_pfxsfx_simulation_gate_128_vs_256_on_long_line(self) -> None:
        ESC = b'\x1b'
        s_ok = '[ПЕРЕДАЧА OK]'.encode('utf-8')
        s_fake = '[ПОДПИСЬ: Поддельная]'.encode('utf-8')
        s_fake_tr = '[ПІДПИС: Підроблений]'.encode('utf-8')
        table = {s_fake + ESC + b'[37;1m': s_fake_tr + ESC + b'[37;1m'}
        line_133 = b'5f6819d80123456789abcdef012345678901.nca : ' + ESC + b'[32;1m' + s_ok + ESC + b'[38;2;255;128;0m' + s_fake + ESC + b'[37;1m'
        self.assertGreater(len(line_133), 128)
        res_128 = simulate_pfxsfx(line_133, table, gate=128)
        self.assertNotIn(s_fake_tr, res_128)
        self.assertEqual(res_128, line_133)
        res_256 = simulate_pfxsfx(line_133, table, gate=256)
        self.assertIn(s_fake_tr, res_256)
        self.assertNotIn(s_fake, res_256)

    def test_compound_status_line_full_simulation(self) -> None:
        ESC = b'\x1b'
        t_ok = '[ПЕРЕДАЧА OK]'.encode('utf-8')
        t_ab = '[ПЕРЕДАЧА ПРЕРВАНА]'.encode('utf-8')
        s_ok = '[ПОДПИСЬ: OK]'.encode('utf-8')
        s_fa = '[ПОДПИСЬ: Поддельная]'.encode('utf-8')
        s_xc = '[ПОДПИСЬ: XCI➡NSP]'.encode('utf-8')

        en_t_ok = '[TRANSFER OK]'.encode('utf-8')
        en_s_ok = '[SIGNATURE: OK]'.encode('utf-8')
        en_s_fa = '[SIGNATURE: Fake]'.encode('utf-8')
        en_s_xc = '[SIGNATURE: XCI➡NSP]'.encode('utf-8')
        en_t_ab = '[TRANSFER INTERRUPTED]'.encode('utf-8')

        table_en = {
            ESC + b'[32;1m' + t_ok + ESC + b'[32;1m' + s_ok + ESC + b'[37;1m': ESC + b'[32;1m' + en_t_ok + ESC + b'[32;1m' + en_s_ok + ESC + b'[37;1m',
            ESC + b'[32;1m' + t_ok + ESC + b'[38;2;255;128;0m' + s_fa + ESC + b'[37;1m': ESC + b'[32;1m' + en_t_ok + ESC + b'[38;2;255;128;0m' + en_s_fa + ESC + b'[37;1m',
            ESC + b'[32;1m' + t_ok + ESC + b'[32;1m' + s_xc + ESC + b'[37;1m': ESC + b'[32;1m' + en_t_ok + ESC + b'[32;1m' + en_s_xc + ESC + b'[37;1m',
            ESC + b'[32;1m' + t_ok + b' ' + ESC + b'[32;1m' + s_ok + ESC + b'[37;1m': ESC + b'[32;1m' + en_t_ok + b' ' + ESC + b'[32;1m' + en_s_ok + ESC + b'[37;1m',
            ESC + b'[32;1m' + t_ok + ESC + b'[37;1m ' + ESC + b'[32;1m' + s_ok + ESC + b'[37;1m': ESC + b'[32;1m' + en_t_ok + ESC + b'[37;1m ' + ESC + b'[32;1m' + en_s_ok + ESC + b'[37;1m',
            ESC + b'[32;1m' + t_ok + ESC + b'[37;1m ' + ESC + b'[38;2;255;128;0m' + s_fa + ESC + b'[37;1m': ESC + b'[32;1m' + en_t_ok + ESC + b'[37;1m ' + ESC + b'[38;2;255;128;0m' + en_s_fa + ESC + b'[37;1m',
            t_ok + b' ' + ESC + b'[32;1m' + s_ok + ESC + b'[37;1m': en_t_ok + b' ' + ESC + b'[32;1m' + en_s_ok + ESC + b'[37;1m',
            t_ok + b' ' + ESC + b'[38;2;255;128;0m' + s_fa + ESC + b'[37;1m': en_t_ok + b' ' + ESC + b'[38;2;255;128;0m' + en_s_fa + ESC + b'[37;1m',
            t_ok + ESC + b'[32;1m' + s_ok + ESC + b'[37;1m': en_t_ok + ESC + b'[32;1m' + en_s_ok + ESC + b'[37;1m',
            t_ok + b' ' + s_ok: en_t_ok + b' ' + en_s_ok,
            ESC + b'[31;1m' + t_ab + ESC + b'[37;1m': ESC + b'[31;1m' + en_t_ab + ESC + b'[37;1m',
        }

        line_compound_ok = b'5f6819d8.nca : ' + ESC + b'[32;1m' + t_ok + ESC + b'[32;1m' + s_ok + ESC + b'[37;1m'
        res_ok = simulate_pfxsfx(line_compound_ok, table_en, gate=256)
        self.assertNotIn(b'\x1b\x1b', res_ok)
        self.assertIn(en_t_ok, res_ok)
        self.assertIn(en_s_ok, res_ok)
        self.assertNotIn(t_ok, res_ok)
        self.assertNotIn(s_ok, res_ok)

        line_compound_fake = b'5f6819d8.nca : ' + ESC + b'[32;1m' + t_ok + ESC + b'[38;2;255;128;0m' + s_fa + ESC + b'[37;1m'
        res_fa = simulate_pfxsfx(line_compound_fake, table_en, gate=256)
        self.assertNotIn(b'\x1b\x1b', res_fa)
        self.assertIn(en_t_ok, res_fa)
        self.assertIn(en_s_fa, res_fa)
        self.assertNotIn(t_ok, res_fa)
        self.assertNotIn(s_fa, res_fa)

        line_spaced = b'5f6819d8.nca : ' + ESC + b'[32;1m' + t_ok + b' ' + ESC + b'[32;1m' + s_ok + ESC + b'[37;1m'
        res_sp = simulate_pfxsfx(line_spaced, table_en, gate=256)
        self.assertIn(en_t_ok, res_sp)
        self.assertIn(en_s_ok, res_sp)
        self.assertNotIn(t_ok, res_sp)
        self.assertNotIn(s_ok, res_sp)

        line_reset_sp = b'5f6819d8.nca : ' + ESC + b'[32;1m' + t_ok + ESC + b'[37;1m ' + ESC + b'[38;2;255;128;0m' + s_fa + ESC + b'[37;1m'
        res_rsp = simulate_pfxsfx(line_reset_sp, table_en, gate=256)
        self.assertIn(en_t_ok, res_rsp)
        self.assertIn(en_s_fa, res_rsp)
        self.assertNotIn(t_ok, res_rsp)
        self.assertNotIn(s_fa, res_rsp)

        line_plain_t_ok = b'5f6819d8.nca : ' + t_ok + b' ' + ESC + b'[32;1m' + s_ok + ESC + b'[37;1m'
        res_pt = simulate_pfxsfx(line_plain_t_ok, table_en, gate=256)
        self.assertIn(en_t_ok, res_pt)
        self.assertIn(en_s_ok, res_pt)
        self.assertNotIn(t_ok, res_pt)
        self.assertNotIn(s_ok, res_pt)

        line_plain_t_fa = b'5f6819d8.nca : ' + t_ok + b' ' + ESC + b'[38;2;255;128;0m' + s_fa + ESC + b'[37;1m'
        res_pt_fa = simulate_pfxsfx(line_plain_t_fa, table_en, gate=256)
        self.assertIn(en_t_ok, res_pt_fa)
        self.assertIn(en_s_fa, res_pt_fa)
        self.assertNotIn(t_ok, res_pt_fa)
        self.assertNotIn(s_fa, res_pt_fa)

        line_plain = b'5f6819d8.nca : ' + t_ok + b' ' + s_ok
        res_plain = simulate_pfxsfx(line_plain, table_en, gate=256)
        self.assertIn(en_t_ok, res_plain)
        self.assertIn(en_s_ok, res_plain)
        self.assertNotIn(t_ok, res_plain)
        self.assertNotIn(s_ok, res_plain)

    def test_patched_nro_contains_256_byte_gate(self) -> None:
        patched_nro = ROOT / 'DBI.905.ru_patched.nro'
        if not patched_nro.is_file():
            self.skipTest('DBI.905.ru_patched.nro not present in repository root')
        from keystone import Ks, KS_ARCH_ARM64, KS_MODE_LITTLE_ENDIAN
        ks = Ks(KS_ARCH_ARM64, KS_MODE_LITTLE_ENDIAN)
        e100, _ = ks.asm('cmp x20, #0x100')
        data = patched_nro.read_bytes()
        pos = data.find(bytes(e100))
        self.assertNotEqual(pos, -1, 'Patched NRO must contain cmp x20, #0x100 instruction')

    def test_validator_accepts_asymmetric_ansi_tokens(self) -> None:
        v = Validator()
        orig = '[32;1mУстановка игры завершена[[ESC]][37;1m'
        trans = '[[ESC]][32;1mВстановлення гри завершено[[ESC]][37;1m'
        self.assertEqual(v.validate_row(orig, trans, 'ua'), [])
        orig_norm = 'Время [[ESC]][32;1m12:00[[ESC]][37;1m'
        trans_bad = 'Час 12:00[[ESC]][37;1m'
        errs_bad = v.validate_row(orig_norm, trans_bad, 'ua')
        self.assertTrue(any('Token mismatch' in e for e in errs_bad))

    def test_dictionary_and_csv_contain_all_11_color_keys(self) -> None:
        wb = openpyxl.load_workbook(self.dict_path, data_only=True)
        ws = wb['Translations']
        wb_keys = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
        for k in self.expected_keys:
            self.assertIn(k, wb_keys, f'Missing in dictionary: {k}')
        for lc in self.languages:
            if lc == 'ru': continue
            csv_path = self.translations_dir / f'{lc}.csv'
            self.assertTrue(csv_path.is_file())
            with open(csv_path, 'r', encoding='utf-8', newline='') as f:
                csv_origs = {r[0] for r in list(csv.reader(f))[1:]}
            for k in self.expected_keys:
                detok_k = k.replace('[[ESC]]', '\\x1b')
                self.assertIn(detok_k, csv_origs, f'Missing {detok_k} in {lc}.csv')

if __name__ == '__main__':
    unittest.main()
