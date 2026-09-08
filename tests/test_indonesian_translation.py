from __future__ import annotations
import csv
import subprocess
import sys
import unittest
from pathlib import Path
import openpyxl

from src.core.text_utils import tokenize, detokenize

ROOT = Path(__file__).resolve().parent.parent
ID_CSV = ROOT / "translations" / "id.csv"
DICT_XLSX = ROOT / "data" / "dictionary.xlsx"
BUILD_SCRIPT = ROOT / "scripts" / "build_translation_bin.py"
OUTPUT_BIN = ROOT / "output" / "translation_id.bin"


class IndonesianTranslationTests(unittest.TestCase):
    def test_id_csv_exists_and_row_count(self) -> None:
        """Verify translations/id.csv exists and has 1294 translation rows."""
        self.assertTrue(ID_CSV.exists(), "translations/id.csv does not exist")
        with open(ID_CSV, "r", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
        self.assertEqual(rows[0], ["original", "translation"], "Invalid header in id.csv")
        self.assertEqual(len(rows) - 1, 1294, f"Expected 1,294 data rows, got {len(rows) - 1}")

    def test_id_csv_completeness(self) -> None:
        """Verify every row in id.csv has non-empty translation."""
        with open(ID_CSV, "r", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))[1:]
        empty_rows = [i for i, r in enumerate(rows, start=2) if not r[1]]
        self.assertEqual(empty_rows, [], f"Empty translations in rows: {empty_rows}")

    def test_dictionary_contains_id_column(self) -> None:
        """Verify data/dictionary.xlsx contains 'id' column and matches id.csv."""
        self.assertTrue(DICT_XLSX.exists(), "data/dictionary.xlsx does not exist")
        wb = openpyxl.load_workbook(DICT_XLSX, read_only=True)
        ws = wb["Translations"]
        headers = [c.value for c in next(ws.iter_rows(max_row=1))]
        self.assertIn("id", headers, "dictionary.xlsx missing 'id' column")
        id_idx = headers.index("id")
        orig_idx = headers.index("Original")

        with open(ID_CSV, "r", encoding="utf-8-sig") as f:
            csv_rows = list(csv.reader(f))[1:]

        excel_rows = list(ws.iter_rows(min_row=2))
        self.assertEqual(len(excel_rows), len(csv_rows), "Row count mismatch between Excel and id.csv")

        mismatches = []
        for i, (e_row, c_row) in enumerate(zip(excel_rows, csv_rows), start=2):
            orig = e_row[orig_idx].value
            trans = e_row[id_idx].value
            if detokenize(str(orig)) != c_row[0]:
                mismatches.append(f"Row {i} original mismatch: {orig!r} vs {c_row[0]!r}")
            if detokenize(str(trans)) != c_row[1]:
                mismatches.append(f"Row {i} translation mismatch: {trans!r} vs {c_row[1]!r}")

        self.assertEqual(mismatches, [], f"Found mismatches: {mismatches[:5]}")

    def test_build_id_binary(self) -> None:
        """Verify scripts/build_translation_bin.py successfully builds translation_id.bin."""
        OUTPUT_BIN.parent.mkdir(parents=True, exist_ok=True)
        result = subprocess.run(
            [sys.executable, str(BUILD_SCRIPT), str(ID_CSV), "-o", str(OUTPUT_BIN)],
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 0, f"Build script failed: {result.stderr}")
        self.assertTrue(OUTPUT_BIN.exists(), "translation_id.bin was not created")
        with open(OUTPUT_BIN, "rb") as f:
            magic = f.read(4)
        self.assertEqual(magic, b"DBIT", "Invalid binary magic")


if __name__ == "__main__":
    unittest.main()
